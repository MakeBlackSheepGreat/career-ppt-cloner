import math

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

import fit

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"

EMU_PER_CM = 360000.0
PT_TO_CM = 0.03528

# 口径档位。presentation 按投影讲解的可读性定，reading 按屏幕自读定，paper 对齐论文画图的 8 磅下限。
PROFILES = ("presentation", "reading", "paper")

ROLE_MIN_PT = {
    "presentation": {
        "display": 36.0, "headline": 24.0, "title": 18.0, "body": 16.0,
        "text": 14.0, "label": 12.0, "backdrop_note": 12.0, "glyph": 12.0, "chrome": 10.0,
    },
    "reading": {
        "display": 24.0, "headline": 16.0, "title": 12.0, "body": 10.0,
        "text": 9.0, "label": 8.0, "backdrop_note": 8.0, "glyph": 8.0, "chrome": 7.0,
    },
    "paper": {
        "display": 8.0, "headline": 8.0, "title": 8.0, "body": 8.0,
        "text": 8.0, "label": 8.0, "backdrop_note": 8.0, "glyph": 8.0, "chrome": 8.0,
    },
}

# 形状超出页面边界的容差，单位厘米。
BOUNDARY_TOLERANCE_CM = 0.05
# 文字框距页面四边的最小留白，单位厘米。
MIN_MARGIN_CM = 0.25
# 上层不透明形状覆盖文字框面积超过这个比例，判为文字被遮挡。
OCCLUSION_RATIO = 0.25
# 正文文字与其底衬色的 WCAG 对比度下限与舒适值。
CONTRAST_MIN = 4.5
CONTRAST_COMFORT = 7.0
# 单页配色数量上限，沿用 Paul Tol 亮色板的 6 色口径。
MAX_PAGE_COLORS = 6
# 单个正文框的最大行数，投影页超过这个行数读者来不及看完。
MAX_BODY_LINES = 6
# 单页可见文字总字数上限，投影页超过这个字数就变成念稿。
MAX_PAGE_CHARS = 220
# 标题字号至少是正文字号的这个倍数，否则层级不分明。
HIERARCHY_RATIO = 1.2
# 同页同角色字号极差上限，单位磅。
SAME_ROLE_SPREAD_PT = 6.0
# 文字框与底衬框中线允许的最大偏差，单位厘米。
CENTRE_TOLERANCE_CM = 0.15
# 行距缺省值，与 fit 模块保持一致。
LINE_HEIGHT_RATIO = 1.22


def qa(tag):
    return "{%s}%s" % (NS_A, tag)


def qp(tag):
    return "{%s}%s" % (NS_P, tag)


def cm_of(value):
    if value is None:
        return None
    return round(Emu(int(value)).cm, 3)


def slide_size_cm(presentation):
    return cm_of(presentation.slide_width), cm_of(presentation.slide_height)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in iter_shapes(shape.shapes):
                yield sub


def bbox_cm(shape):
    if None in (shape.left, shape.top, shape.width, shape.height):
        return None
    return (cm_of(shape.left), cm_of(shape.top), cm_of(shape.width), cm_of(shape.height))


def overlap_area_cm(a, b):
    left = max(a[0], b[0])
    top = max(a[1], b[1])
    right = min(a[0] + a[2], b[0] + b[2])
    bottom = min(a[1] + a[3], b[1] + b[3])
    if right <= left or bottom <= top:
        return 0.0
    return (right - left) * (bottom - top)


def centre_cm(box):
    return box[0] + box[2] / 2.0, box[1] + box[3] / 2.0


def role_of(text, width, height, slide_w, slide_h, font_size, char_count):
    if char_count <= 1:
        return "glyph"
    if width and height and width >= slide_w * 0.88 and height >= slide_h * 0.55:
        return "backdrop_note"
    if font_size and font_size >= 44:
        return "display"
    if font_size and font_size >= 30:
        return "headline"
    if font_size and font_size >= 22:
        return "title"
    if char_count <= 8 and (font_size or 0) >= 18:
        return "label"
    if char_count >= 55:
        return "body"
    return "text"


def is_chrome(slot):
    top = slot.get("top") or 0.0
    if top >= 3.2:
        return False
    return slot.get("chars", 0) <= 14 and slot.get("role") in ("text", "label")


def body_properties(shape):
    frame = shape.text_frame
    body = frame._txBody.find(qa("bodyPr"))
    result = {
        "insets": [0.254, 0.254, 0.127, 0.127],
        "anchor": "TOP",
        "autofit": "none",
        "font_scale": 1.0,
        "line_spacing": None,
    }
    if body is None:
        return result
    for key, attribute in (("l_ins", "lIns"), ("r_ins", "rIns"), ("t_ins", "tIns"), ("b_ins", "bIns")):
        raw = body.get(attribute)
        if raw is not None:
            index = {"l_ins": 0, "r_ins": 1, "t_ins": 2, "b_ins": 3}[key]
            result["insets"][index] = round(Emu(int(raw)).cm, 3)
    anchor = body.get("anchor")
    if anchor:
        result["anchor"] = anchor
    for child in body:
        tag = etree.QName(child).localname
        if tag in ("normAutofit", "spAutoFit", "noAutofit"):
            result["autofit"] = tag
            if tag == "normAutofit" and child.get("fontScale"):
                result["font_scale"] = int(child.get("fontScale")) / 100000.0
    paragraph = frame.paragraphs[0] if frame.paragraphs else None
    if paragraph is not None:
        ppr = paragraph._p.find(qa("pPr"))
        if ppr is not None:
            spacing = ppr.find(qa("lnSpc"))
            if spacing is not None:
                for child in spacing:
                    tag = etree.QName(child).localname
                    if tag == "spcPct":
                        result["line_spacing"] = round(int(child.get("val")) / 100000.0, 3)
                    elif tag == "spcPts":
                        result["line_spacing"] = round(int(child.get("val")) / 100.0, 2)
    return result


def usable_area_cm(slot):
    insets = slot.get("insets") or [0.254, 0.254, 0.127, 0.127]
    left, right, top, bottom = insets
    width = max(0.4, (slot.get("width") or 0.0) - left - right)
    height = max(0.3, (slot.get("height") or 0.0) - top - bottom)
    return width, height


def text_runs(shape):
    runs = []
    for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = run.font.size.pt if run.font.size else None
            runs.append({
                "paragraph": paragraph_index,
                "text": run.text,
                "size_pt": size,
                "bold": bool(run.font.bold) if run.font.bold is not None else False,
                "font_name": run.font.name,
                "color": color_of_font(run.font),
            })
    return runs


def color_of_font(font):
    color = font.color
    if color is None or color.type is None:
        return None
    kind = str(color.type)
    if "RGB" in kind:
        return "#" + str(color.rgb)
    if "THEME" in kind:
        return "theme:" + str(color.theme_color)
    return None


def theme_scheme(presentation):
    mapping = {}
    package = presentation.part.package
    theme_part = None
    for part in package.iter_parts():
        name = str(part.partname)
        if "theme" in name and name.endswith(".xml"):
            theme_part = part
            break
    if theme_part is None:
        return mapping
    root = etree.fromstring(theme_part.blob)
    scheme = root.find(".//" + qa("clrScheme"))
    if scheme is None:
        return mapping
    for child in scheme:
        slot = etree.QName(child).localname
        for color_child in child:
            tag = etree.QName(color_child).localname
            if tag == "srgbClr":
                value = color_child.get("val")
                if value:
                    mapping[slot] = "#" + value.upper()
            elif tag == "sysClr":
                value = color_child.get("lastClr") or color_child.get("val")
                if value and len(value) == 6:
                    mapping[slot] = "#" + value.upper()
    return mapping


def resolve_color(value, scheme):
    if not value:
        return None
    if value.startswith("#"):
        return value.upper()
    if value.startswith("theme:"):
        key = value.split(":", 1)[1]
        aliases = {
            "BACKGROUND_1": "bg1", "BACKGROUND_2": "bg2",
            "TEXT_1": "tx1", "TEXT_2": "tx2",
            "ACCENT_1": "accent1", "ACCENT_2": "accent2", "ACCENT_3": "accent3",
            "ACCENT_4": "accent4", "ACCENT_5": "accent5", "ACCENT_6": "accent6",
            "HYPERLINK": "hlink", "FOLLOWED_HYPERLINK": "folHlink",
        }
        key = aliases.get(key, key)
        return scheme.get(key)
    return None


def shape_fill(shape):
    element = shape._element
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return {"kind": "picture", "color": None}
    sp_pr = element.find(qp("spPr"))
    if sp_pr is None:
        sp_pr = element.find(qa("spPr"))
    if sp_pr is None:
        return {"kind": "none", "color": None}
    solid = sp_pr.find(qa("solidFill"))
    if solid is None:
        if sp_pr.find(qa("gradFill")) is not None:
            return {"kind": "gradient", "color": None}
        if sp_pr.find(qa("pattFill")) is not None:
            return {"kind": "pattern", "color": None}
        if sp_pr.find(qa("noFill")) is not None:
            return {"kind": "none", "color": None}
        return {"kind": "none", "color": None}
    for child in solid:
        tag = etree.QName(child).localname
        value = child.get("val")
        if tag == "srgbClr" and value:
            return {"kind": "solid", "color": "#" + value.upper()}
        if tag == "schemeClr" and value:
            return {"kind": "solid", "color": "theme:" + value}
    return {"kind": "solid", "color": None}


def slide_background(slide, presentation, scheme):
    chain = [slide, slide.slide_layout, slide.slide_layout.slide_master]
    for item in chain:
        element = item._element
        background = element.find(qp("cSld") + "/" + qp("bg"))
        if background is None:
            background = element.find(qp("bg"))
        if background is None:
            continue
        bg_pr = background.find(qp("bgPr"))
        if bg_pr is None:
            bg_ref = background.find(qp("bgRef"))
            if bg_ref is not None:
                color = bg_ref.find(qa("schemeClr"))
                if color is not None and color.get("val"):
                    return resolve_color("theme:" + color.get("val"), scheme)
            continue
        for child in bg_pr:
            tag = etree.QName(child).localname
            if tag == "solidFill":
                for color_child in child:
                    value = color_child.get("val")
                    if value:
                        return resolve_color("theme:" + value if etree.QName(color_child).localname == "schemeClr" else "#" + value, scheme)
            if tag == "blipFill":
                return "picture"
            if tag == "gradFill":
                return "gradient"
    return "#FFFFFF"


def hex_to_rgb(hex_color):
    value = hex_color.lstrip("#")
    return tuple(int(value[index:index + 2], 16) / 255.0 for index in (0, 2, 4))


def relative_luminance(hex_color):
    def channel(value):
        return value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4
    red, green, blue = hex_to_rgb(hex_color)
    return 0.2126 * channel(red) + 0.7152 * channel(green) + 0.0722 * channel(blue)


def contrast_ratio(foreground, background):
    if not foreground or not background:
        return None
    if foreground == "picture" or background == "picture":
        return None
    if foreground == "gradient" or background == "gradient":
        return None
    first = relative_luminance(foreground)
    second = relative_luminance(background)
    lighter = max(first, second)
    darker = min(first, second)
    return round((lighter + 0.05) / (darker + 0.05), 2)


def describe_slots(slide, slide_w, slide_h, scheme):
    slots = []
    for shape in iter_shapes(slide.shapes):
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        box = bbox_cm(shape)
        if box is None:
            continue
        runs = text_runs(shape)
        first = runs[0] if runs else {}
        props = body_properties(shape)
        slot = {
            "slot": len(slots) + 1,
            "shape_id": shape.shape_id,
            "name": shape.name,
            "text": text,
            "chars": len(text),
            "left": box[0], "top": box[1], "width": box[2], "height": box[3],
            "insets": props["insets"],
            "line_spacing": props["line_spacing"],
            "anchor": props["anchor"],
            "autofit": props["autofit"],
            "font_scale": props["font_scale"],
            "paragraphs": len(shape.text_frame.paragraphs),
            "font_size": first.get("size_pt"),
            "font_name": first.get("font_name"),
            "bold": first.get("bold"),
            "color": first.get("color"),
            "runs": runs,
        }
        slot["role"] = role_of(text, box[2], box[3], slide_w, slide_h,
                               slot["font_size"], slot["chars"])
        slot["chrome"] = is_chrome(slot)
        slots.append(slot)
    return slots


def find_backplate(slot, shapes, slide_area):
    subject = (slot["left"], slot["top"], slot["width"], slot["height"])
    if subject[2] <= 0 or subject[3] <= 0:
        return None
    best = None
    best_score = -1e18
    for other in shapes:
        if other is None:
            continue
        if other.shape_type == MSO_SHAPE_TYPE.PICTURE:
            continue
        if other.has_text_frame and other.text_frame.text.strip():
            continue
        box = bbox_cm(other)
        if box is None or box[2] <= 0 or box[3] <= 0:
            continue
        if box[2] * box[3] >= slide_area * 0.62:
            continue
        cx, cy = centre_cm(subject)
        if not (box[0] <= cx <= box[0] + box[2] and box[1] <= cy <= box[1] + box[3]):
            continue
        covered = overlap_area_cm(box, subject) / float(subject[2] * subject[3])
        if covered < 0.90:
            continue
        size_ratio = (box[2] * box[3]) / float(subject[2] * subject[3])
        if size_ratio > 9.0 or size_ratio < 0.95:
            continue
        offset = (abs(centre_cm(box)[0] - cx) / float(subject[2])
                  + abs(centre_cm(box)[1] - cy) / float(subject[3]))
        score = covered * 2.0 - abs(math.log(size_ratio)) - offset
        if score > best_score:
            best_score = score
            best = other
    return best


def open_presentation(path):
    presentation = Presentation(path)
    slide_w, slide_h = slide_size_cm(presentation)
    scheme = theme_scheme(presentation)
    return presentation, slide_w, slide_h, scheme


def profile_limits(profile):
    if profile not in ROLE_MIN_PT:
        raise SystemExit("未知口径档位：%s（可选 %s）" % (profile, "、".join(PROFILES)))
    return ROLE_MIN_PT[profile]
